import os
from openpyxl import load_workbook, Workbook

class ExcelDataParser:

    def __init__(self, excel_file_path):
        self.excel_file_path = excel_file_path
        if not os.path.exists(self.excel_file_path):
            raise FileNotFoundError(f"The file '{self.excel_file_path}' does not exist.")
        self.workbook = load_workbook(self.excel_file_path)

    def create_excel_file(self, sheet_name="Sheet1"):
        try:
            self.workbook = load_workbook(self.excel_file_path)
        except FileNotFoundError:
            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)
            self.workbook.create_sheet(sheet_name)
            self.save_workbook()

    def create_excel_sheet(self, sheet_name):
        if sheet_name not in self.workbook.sheetnames:
            self.workbook.create_sheet(sheet_name)
        self.save_workbook()

    def save_workbook(self):
        self.workbook.save(self.excel_file_path)

    def apply_excel_template(self, sheet_name, identifiers):
        if sheet_name not in self.workbook.sheetnames:
            self.create_excel_sheet(sheet_name)

        headers = self.get_column_names(sheet_name)

        for i, identifier in enumerate(identifiers):
            start_col_index = i + 1
            if identifier not in headers:
                self.write_in_cell(sheet_name, row=1, column=start_col_index, cell_value=identifier)

        self.save_workbook()

    def get_column_values(self, sheet_name, column):
        sheet = self.workbook[sheet_name]
        return [sheet[f"{column}{row}"].value for row in range(1, sheet.max_row + 1)]

    def get_column_names(self, sheet_name):
        sheet = self.workbook[sheet_name]
        return [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=sheet.max_row))]

    def write_in_cell(self, sheet_name, row, column, cell_value):
        sheet = self.workbook[sheet_name]
        sheet.cell(row=row, column=column, value=cell_value)

    def rename_column(self, sheet_name, old_column_name, new_column_name):
        headers = self.get_column_names(sheet_name)
        if old_column_name not in headers:
            raise ValueError(f"Column '{old_column_name}' does not exist.")
        column_index = headers.index(old_column_name) + 1
        self.write_in_cell(sheet_name, row=1, column=column_index, cell_value=new_column_name)

    def add_new_column(self, sheet_name, column_name, default_value=None):
        headers = self.get_column_names(sheet_name)
        if column_name in headers:
            print(f"Column '{column_name}' already exists.")
            return
        new_col_index = len(headers) + 1
        self.write_in_cell(sheet_name, row=1, column=new_col_index, cell_value=column_name)
        sheet = self.workbook[sheet_name]
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=new_col_index, value=default_value)

        self.save_workbook()

    def fill_excel_with_data(self, sheet_name, data):

        # Check if the sheet exists, create it if not
        if sheet_name not in self.workbook.sheetnames:
            self.create_excel_sheet(sheet_name)

        # Retrieve all headers
        headers = self.get_column_names(sheet_name)

        sheet = self.workbook[sheet_name]

        # Ensure all keys in data exist as headers
        for column_name in data.keys():
            if column_name not in headers:
                self.add_new_column(sheet_name, column_name)
                headers = self.get_column_names(sheet_name)  # Update headers after adding a new column

        # Write data to the next available row in each column
        for column_name, value in data.items():
            column_index = headers.index(column_name) + 1  # Convert header index to 1-based column index

            # Find the next empty cell in the column
            for row in range(2, sheet.max_row + 2):  # Start at row 2 to skip headers
                if sheet.cell(row=row, column=column_index).value is None:
                    self.write_in_cell(sheet_name, row=row, column=column_index, cell_value=value)
                    break

        # Save the workbook after writing data
        self.save_workbook()


